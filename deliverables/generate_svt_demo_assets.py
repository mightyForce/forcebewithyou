from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent
SLIDES_DIR = ROOT / "svt_demo_slides"
PPTX_PATH = ROOT / "SVT-Advisor-Copilot-Demo.pptx"
VIDEO_LIST = ROOT / "svt_video_concat.txt"

W, H = 1920, 1080
NAVY = "#083B72"
BLUE = "#0B67B2"
TEAL = "#00A3A3"
PURPLE = "#6E56CF"
GREEN = "#1B8A5A"
ORANGE = "#F59E0B"
RED = "#C2413B"
INK = "#112A46"
MUTED = "#56708B"
BG = "#F6F8FC"
WHITE = "#FFFFFF"
PALE_BLUE = "#E7F1FA"
PALE_GREEN = "#E9F6EF"
PALE_PURPLE = "#F0ECFF"
PALE_ORANGE = "#FFF4DE"

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def draw_wrapped(draw, xy, text, width, size=28, color=INK, bold=False, spacing=9):
    x, y = xy
    chars = max(12, int(width / (size * 0.54)))
    for line in wrap(text, width=chars):
        draw.text((x, y), line, font=font(size, bold), fill=color)
        y += size + spacing
    return y


def rounded(draw, box, fill=WHITE, outline=None, radius=28, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, color=BLUE, width=8):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = end
    draw.polygon([(x1, y1), (x1 - 22, y1 - 13), (x1 - 22, y1 + 13)], fill=color)


def base(slide_no, kicker="SVT ADVISOR COPILOT"):
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, W, 92), fill=NAVY)
    draw.text((80, 26), kicker, font=font(28, True), fill=WHITE)
    draw.text((W - 245, 30), "FICTIONAL DEMO", font=font(20, True), fill="#B7D8F5")
    draw.text((80, H - 50), "Salesforce Data Cloud + Agentforce for Sales", font=font(18), fill=MUTED)
    draw.text((W - 120, H - 50), f"{slide_no:02d}", font=font(20, True), fill=MUTED)
    return img, draw


def title(draw, heading, subheading=None):
    draw.text((80, 150), heading, font=font(54, True), fill=INK)
    if subheading:
        draw_wrapped(draw, (82, 225), subheading, 1150, size=25, color=MUTED)


def pill(draw, box, text, fill, text_color=WHITE):
    rounded(draw, box, fill=fill, radius=22)
    bx0, by0, bx1, by1 = box
    tw = draw.textbbox((0, 0), text, font=font(20, True))[2]
    draw.text((bx0 + (bx1 - bx0 - tw) / 2, by0 + 13), text, font=font(20, True), fill=text_color)


def metric_card(draw, box, label, value, accent=BLUE):
    rounded(draw, box, fill=WHITE, outline="#D9E4EF")
    x0, y0, x1, y1 = box
    draw.rectangle((x0, y0, x0 + 10, y1), fill=accent)
    draw.text((x0 + 34, y0 + 28), label, font=font(20, True), fill=MUTED)
    draw.text((x0 + 34, y0 + 72), value, font=font(38, True), fill=INK)


def slide_cover():
    img, draw = base(1, "SVT MOTORS")
    draw.ellipse((1190, 160, 1790, 760), fill=PALE_BLUE)
    draw.ellipse((1330, 280, 1620, 570), fill=PALE_PURPLE)
    title(draw, "SVT Advisor Copilot", "A five-minute Salesforce demo: unified rider data, governed AI guidance, and human-approved action.")
    pill(draw, (80, 370, 380, 425), "DATA CLOUD", TEAL)
    pill(draw, (400, 370, 760, 425), "AGENTFORCE FOR SALES", PURPLE)
    pill(draw, (780, 370, 1050, 425), "SALES APP", BLUE)
    for cx, cy, label, color in [(1280, 300, "DATA", TEAL), (1530, 470, "AI", PURPLE), (1280, 650, "ADVISOR", BLUE)]:
        draw.ellipse((cx - 95, cy - 95, cx + 95, cy + 95), fill=color)
        tw = draw.textbbox((0, 0), label, font=font(22, True))[2]
        draw.text((cx - tw / 2, cy - 12), label, font=font(22, True), fill=WHITE)
    arrow(draw, (1370, 340), (1460, 430), WHITE, 6)
    arrow(draw, (1460, 540), (1370, 625), WHITE, 6)
    draw.text((80, 800), "Built with synthetic data only. SVT Motors is a fictional demonstration company.", font=font(24), fill=INK)
    return img


def slide_business_moment():
    img, draw = base(2)
    title(draw, "The advisor’s business moment", "Turn fragmented rider signals into a governed, relevant follow-up conversation.")
    cards = [
        ("CRM", "Contact, Tasks", PALE_BLUE, BLUE),
        ("DEALER / DMS", "Vehicle + service", PALE_GREEN, GREEN),
        ("DIGITAL", "Model page + configurator", PALE_PURPLE, PURPLE),
    ]
    x = 90
    for label, text, bg, color in cards:
        rounded(draw, (x, 390, x + 410, 610), fill=bg, outline=color)
        draw.text((x + 35, 430), label, font=font(28, True), fill=color)
        draw_wrapped(draw, (x + 35, 490), text, 310, size=25, color=INK)
        x += 490
    arrow(draw, (500, 500), (565, 500), BLUE)
    arrow(draw, (990, 500), (1055, 500), BLUE)
    rounded(draw, (360, 720, 1560, 895), fill=WHITE, outline="#D9E4EF")
    draw.text((420, 760), "Advisor question", font=font(24, True), fill=MUTED)
    draw.text((420, 808), "“Who should I follow up with, and what can I say with confidence?”", font=font(33, True), fill=INK)
    return img


def slide_architecture():
    img, draw = base(3)
    title(draw, "Unified data foundation", "Data Cloud harmonizes the signals; the agent uses Flows for governed retrieval and action.")
    nodes = [
        (100, 390, 330, 555, "CRM", "riders + tasks", BLUE),
        (100, 640, 330, 805, "DEALER / DMS", "service history", GREEN),
        (100, 890, 330, 1050, "DIGITAL", "engagement events", PURPLE),
        (610, 505, 950, 755, "DATA CLOUD", "identity + readiness", TEAL),
        (1220, 505, 1545, 755, "AGENTFORCE", "retrieve + guardrails", PURPLE),
        (1640, 555, 1840, 705, "SALES", "advisor action", BLUE),
    ]
    for x0, y0, x1, y1, label, text, color in nodes:
        rounded(draw, (x0, y0, x1, y1), fill=WHITE, outline=color, width=4)
        draw.text((x0 + 25, y0 + 30), label, font=font(23, True), fill=color)
        draw_wrapped(draw, (x0 + 25, y0 + 80), text, x1 - x0 - 45, size=20, color=INK)
    for sy in (470, 720, 970):
        arrow(draw, (340, sy), (580, 625), TEAL, 6)
    arrow(draw, (965, 630), (1190, 630), PURPLE, 8)
    arrow(draw, (1555, 630), (1610, 630), BLUE, 8)
    return img


def slide_readiness():
    img, draw = base(4)
    title(draw, "Transparent readiness logic", "A simple demo classification—not a prediction, offer, or approval.")
    metric_card(draw, (110, 360, 560, 535), "LATEST ODOMETER", "10,400 km", BLUE)
    metric_card(draw, (620, 360, 1070, 535), "ENGAGEMENT EVENTS", "2", TEAL)
    metric_card(draw, (1130, 360, 1580, 535), "DEMO LABEL", "HIGH", GREEN)
    draw.text((130, 635), "Rule 1", font=font(22, True), fill=BLUE)
    draw.text((260, 635), "Odometer ≥ 10,000 km AND at least one engagement → HIGH", font=font(25), fill=INK)
    draw.text((130, 705), "Rule 2", font=font(22, True), fill=PURPLE)
    draw.text((260, 705), "Either signal is present → MEDIUM", font=font(25), fill=INK)
    draw.text((130, 775), "Fallback", font=font(22, True), fill=MUTED)
    draw.text((260, 775), "Neither signal is present → LOW", font=font(25), fill=INK)
    rounded(draw, (120, 865, 1700, 970), fill=PALE_ORANGE, outline=ORANGE)
    draw_wrapped(draw, (155, 895), "Important: HIGH means the demo criteria were met. It never means a rider is eligible for a discount, financing, warranty, or an automatic message.", 1480, size=21, color=INK)
    return img


def slide_console():
    img, draw = base(5)
    title(draw, "SVT Advisor Console in Sales", "A single advisor workspace for the Contact, tasks, and Employee Agent.")
    rounded(draw, (90, 330, 1830, 950), fill=WHITE, outline="#CEDAE7", width=3)
    draw.rectangle((90, 330, 1830, 410), fill="#EEF4FA")
    draw.text((125, 355), "Ananya Rao", font=font(28, True), fill=INK)
    draw.text((125, 450), "Contact", font=font(20, True), fill=MUTED)
    for idx, (k, v) in enumerate([("Email", "ananya.rao@example.test"), ("Vehicle", "SVT Stride 200 Demo"), ("Contact owner", "Sanjay Chari")]):
        draw.text((125, 505 + idx * 70), k, font=font(19, True), fill=MUTED)
        draw.text((340, 505 + idx * 70), v, font=font(21), fill=INK)
    rounded(draw, (835, 445, 1260, 875), fill=PALE_GREEN, outline=GREEN)
    draw.text((875, 480), "Activities", font=font(23, True), fill=GREEN)
    draw.text((875, 545), "• SVT Demo: Advisor Follow-up", font=font(19), fill=INK)
    draw.text((875, 590), "  Assigned to Sanjay Chari", font=font(17), fill=MUTED)
    draw.text((875, 680), "• No customer message sent", font=font(19), fill=INK)
    rounded(draw, (1320, 445, 1780, 875), fill=PALE_PURPLE, outline=PURPLE)
    draw.text((1360, 480), "SVT Advisor Copilot", font=font(22, True), fill=PURPLE)
    draw_wrapped(draw, (1360, 550), "“Find Ananya and explain readiness.”", 360, size=18, color=INK)
    draw_wrapped(draw, (1360, 655), "HIGH • ModelPageViewed • grounded service context", 360, size=18, color=INK)
    return img


def slide_guardrails():
    img, draw = base(6)
    title(draw, "Guardrails by design", "The agent retrieves facts and follows approval boundaries.")
    items = [
        ("Exact identity match", "Lookup is based on a rider ID or exact synthetic email.", GREEN),
        ("No commercial approvals", "Discount, pricing, financing, warranty, and payment approval are refused.", RED),
        ("No mechanical advice", "The agent does not diagnose faults or provide safety advice.", ORANGE),
        ("Human confirmation", "A follow-up Task is created only after an explicit confirmation.", PURPLE),
    ]
    y = 330
    for headline, text, color in items:
        rounded(draw, (150, y, 1770, y + 135), fill=WHITE, outline="#D6E2ED")
        draw.ellipse((185, y + 35, 245, y + 95), fill=color)
        draw.text((285, y + 28), headline, font=font(25, True), fill=INK)
        draw_wrapped(draw, (285, y + 70), text, 1360, size=19, color=MUTED)
        y += 155
    return img


def slide_action():
    img, draw = base(7)
    title(draw, "Human-approved action", "The only write action creates an internal Salesforce Task.")
    steps = [
        ("1", "Advisor asks", "Create a follow-up Task"),
        ("2", "Agent confirms", "No customer message is sent"),
        ("3", "Advisor says yes", "Explicit confirmation required"),
        ("4", "Flow creates Task", "Assigned to human advisor"),
    ]
    x = 100
    for no, headline, text in steps:
        rounded(draw, (x, 430, x + 370, 700), fill=WHITE, outline="#D9E4EF")
        draw.ellipse((x + 35, 460, x + 100, 525), fill=BLUE)
        draw.text((x + 57, 476), no, font=font(24, True), fill=WHITE)
        draw_wrapped(draw, (x + 35, 555), headline, 300, size=24, color=INK, bold=True)
        draw_wrapped(draw, (x + 35, 620), text, 300, size=18, color=MUTED)
        if x < 1300:
            arrow(draw, (x + 380, 565), (x + 450, 565), TEAL, 6)
        x += 455
    rounded(draw, (260, 800, 1660, 915), fill=PALE_GREEN, outline=GREEN)
    draw.text((315, 842), "Outcome: auditable Task, assigned to the advisor, no automatic outbound contact.", font=font(25, True), fill=GREEN)
    return img


def slide_live_sequence():
    img, draw = base(8)
    title(draw, "Five-minute live demo sequence", "Use the Sales app, dashboard, and Employee Agent to demonstrate the whole loop.")
    steps = [
        ("Dashboard", "Show engagement and odometer signals."),
        ("Lookup", "Find ananya.rao@example.test."),
        ("Readiness", "Explain HIGH + ModelPageViewed as demo logic."),
        ("Guardrail", "Ask for discount + financing approval."),
        ("Task", "Request and explicitly confirm internal follow-up."),
    ]
    y = 330
    for idx, (headline, text) in enumerate(steps, 1):
        draw.ellipse((140, y - 5, 205, y + 60), fill=PURPLE)
        draw.text((162, y + 12), str(idx), font=font(23, True), fill=WHITE)
        draw.text((250, y), headline, font=font(25, True), fill=INK)
        draw.text((510, y), text, font=font(22), fill=MUTED)
        y += 105
    rounded(draw, (120, 860, 1780, 960), fill=PALE_ORANGE, outline=ORANGE)
    draw.text((160, 895), "Always state: SVT is fictional and all demo data is synthetic.", font=font(23, True), fill=INK)
    return img


def slide_close():
    img, draw = base(9)
    title(draw, "What the demo proves", "A governed pattern for turning unified data into an advisor-led next step.")
    outcomes = [
        ("Unify", "Data Cloud brings rider, vehicle, service, and engagement signals together.", TEAL),
        ("Guide", "Agentforce retrieves facts and explains transparent demo logic.", PURPLE),
        ("Protect", "Guardrails block commercial approval and require human confirmation.", RED),
        ("Act", "A Flow creates an auditable internal Task assigned to an advisor.", GREEN),
    ]
    x = 95
    for title_text, body, color in outcomes:
        rounded(draw, (x, 400, x + 400, 710), fill=WHITE, outline=color, width=3)
        draw.text((x + 35, 445), title_text, font=font(31, True), fill=color)
        draw_wrapped(draw, (x + 35, 520), body, 320, size=21, color=INK)
        x += 440
    draw.text((95, 835), "Next: production data governance, consent model, dealer ownership rules, and deployment automation.", font=font(25), fill=MUTED)
    return img


def main():
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    creators = [
        slide_cover, slide_business_moment, slide_architecture, slide_readiness,
        slide_console, slide_guardrails, slide_action, slide_live_sequence, slide_close,
    ]
    slide_paths = []
    for idx, create in enumerate(creators, 1):
        path = SLIDES_DIR / f"slide_{idx:02d}.png"
        create().save(path, quality=95)
        slide_paths.append(path)

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in slide_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX_PATH)

    with VIDEO_LIST.open("w", encoding="utf-8") as f:
        for path in slide_paths:
            f.write(f"file '{path.as_posix()}'\n")
            f.write("duration 7\n")
        f.write(f"file '{slide_paths[-1].as_posix()}'\n")


if __name__ == "__main__":
    main()
