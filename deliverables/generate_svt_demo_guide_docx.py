#!/usr/bin/env python3
"""Generate SVT-Advisor-Copilot-Demo-Guide.docx from repo markdown sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "SVT-Advisor-Copilot-Demo-Guide.docx"

LEAD_INSTRUCTIONS = """You help dealer advisors qualify prospects and schedule test rides.

When the advisor asks about a prospect or lead:
Run SVT Get Lead Context V2 using the prospect's email or name as searchText.
If no lead is returned, ask the advisor for an email address or phone number, then retry SVT Get Lead Context V2.
Run SVT Get Lead Intent v2 using prospectId from SVT Get Lead Context V2.
Explain the intent score (HIGH, MEDIUM, or LOW) and which engagement events drove it.
Run SVT Recommend Dealer with these three inputs:
city from SVT Get Lead Context V2
pinCode from SVT Get Lead Context V2
model from preferredModel returned by SVT Get Lead Intent v2
The model input is mandatory. Never run SVT Recommend Dealer without it.
Recommend the dealer and explain the match (PIN match or city match).

If hasOpenTestRideTask from SVT Get Lead Context V2 is true:
Tell the advisor an open test ride task already exists for this prospect.
Include openTaskSubject from SVT Get Lead Context V2 in your response.
Do NOT ask "Should I create a test ride task?"
Do NOT offer or suggest creating a new test ride task.
Do NOT run SVT Create Lead Test Ride Task under any circumstances.
If the advisor explicitly asks to create a task, explain that an open task already exists and no duplicate will be created.

If contactPermission from SVT Get Lead Context V2 is false:
Still complete intent and dealer steps. Share intent score and recommended dealer with the advisor.
Clearly state that the prospect has not given contact consent.
Do NOT ask to create a test ride task.
Do NOT run SVT Create Lead Test Ride Task under any circumstances.

If contactPermission is true AND hasOpenTestRideTask is false:
Ask the advisor: "Should I create a test ride task for this prospect?"
Only run SVT Create Lead Test Ride Task after the advisor explicitly confirms.
Pass userConfirmed=true only after confirmation.
Pass leadId from SVT Get Lead Context V2 (CRM Lead Id). Never ask the advisor for a Salesforce Id.
Pass dealerName from SVT Recommend Dealer and model from preferredModel in SVT Get Lead Intent v2.
Pass contactPermission from SVT Get Lead Context V2.

If SVT Create Lead Test Ride Task returns taskStatus ALREADY_EXISTS:
Tell the advisor an open test ride task already exists (include the message from the action).
Do not offer or attempt to create another.

If taskStatus is CREATED:
Confirm once that the task was created and assigned to the advisor.
If emailSent is true, tell the advisor a confirmation email was sent to the prospect.
If emailSent is false, tell the advisor the task was created but no confirmation email was sent.

Never run SVT Create Lead Test Ride Task if the advisor has not confirmed, contactPermission is false, or hasOpenTestRideTask is true.

DOCUMENT GROUNDING (SVT Advisor Quick Guides — agent-level Data Library):
The agent has Data Library "SVT Advisor Quick Guides" attached at agent level under Data → Data Library.
Use Answer Question with Knowledge for document lookups.
Use Flow actions for intent, dealer, consent, and task creation. NEVER infer these from PDFs.
Use the Data Library for test ride process, model talking points, key specs, and objection scripts.
Always cite the document title. If not found: "I don't have that in the approved SVT documents."
Do not invent specs, warranty terms, or pricing not stated in retrieved documents.

GUARDRAILS:
Do not approve discounts, financing, or pricing.
Do not provide medical or legal advice.
All prospect names and data are synthetic demo data."""


def set_update_fields(doc: Document):
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_title_page(doc: Document):
    p = doc.add_paragraph()
    p.style = "Title"
    p.add_run("SVT Advisor Copilot Demo Guide")
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run("SVT Motors Data Cloud + Agentforce (fictional demo)\n").bold = True
    meta.add_run(
        "Version: August 2026\n"
        "Includes: Rider retention, Lead engagement, Streaming Ingestion API, "
        "Agentforce Data Library (Path A)\n"
        "Out of scope: Data Cloud unstructured RAG (Path B / S3 / UDLO)\n"
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_heading(doc: Document, text: str, level: int):
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, style: str | None = None):
    if not text.strip():
        return
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        else:
            p.add_run(part)


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 * level)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        else:
            p.add_run(part)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            table.rows[ri].cells[ci].text = cell.strip()
    doc.add_paragraph()


def parse_md_table(lines: list[str]) -> list[list[str]] | None:
    if len(lines) < 2 or "|" not in lines[0]:
        return None
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            break
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows if rows else None


def skip_mermaid_block(lines: list[str], i: int) -> int:
    if lines[i].strip().startswith("```"):
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            i += 1
        return i + 1
    return i


def md_to_docx(doc: Document, md_path: Path, skip_sections: set[str] | None = None):
    skip_sections = skip_sections or set()
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i = skip_mermaid_block(lines, i)
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped.lstrip("#").strip()
            if title in skip_sections:
                i += 1
                while i < len(lines) and not (
                    lines[i].strip().startswith("## ") and not lines[i].strip().startswith("###")
                ):
                    if lines[i].strip().startswith("#") and lines[i].count("#") <= level:
                        break
                    i += 1
                continue
            doc_level = min(level, 4)
            add_heading(doc, title, doc_level)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_md_table(table_lines)
            if rows:
                add_table(doc, rows)
            continue

        if stripped.startswith(">"):
            quote = stripped.lstrip(">").strip().strip('"')
            add_para(doc, quote)
            i += 1
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.add_run(m.group(2))
            i += 1
            continue

        if stripped.startswith("[") and "](http" in stripped:
            add_para(doc, stripped)
            i += 1
            continue

        if stripped:
            add_para(doc, stripped)
        i += 1


def add_supplement_sections(doc: Document):
    doc.add_page_break()
    add_heading(doc, "14. Document grounding — Agentforce Data Library (Path A)", 1)
    add_para(
        doc,
        "Path B (S3, UDLO, Search Index, Retriever) is dropped. Use Agentforce Data Library only for advisor PDFs.",
    )

    add_heading(doc, "14.1 PDF set — SVT Advisor Quick Guides", 2)
    add_table(
        doc,
        [
            ["ID", "File", "Purpose"],
            ["A1", "SVT_Advisor_Test_Ride_Playbook.pdf", "Test ride steps"],
            ["A2", "SVT_Stride_200_Advisor_OnePager.pdf", "Stride 200 features + key specs (3.2 kWh, motor 4.2 kW)"],
            ["A3", "SVT_Urban_125_Advisor_OnePager.pdf", "Urban 125 positioning"],
            ["A4", "SVT_Objection_Handling_Guide.pdf", "Objection handling (no rupee amounts)"],
        ],
    )

    add_heading(doc, "14.2 Org setup (summary)", 2)
    for step in [
        "Setup → Agentforce Data Library → New → SVT Advisor Quick Guides",
        "Upload all 4 PDFs; wait until each file shows Ready",
        "Agent Builder → SVT Advisor Copilot → Data → Data Library → select library",
        "Lead Engagement subagent → add Answer Question with Knowledge action",
        "Paste subagent instructions (Appendix A in this document)",
        "Save → Commit Version → Activate",
    ]:
        add_bullet(doc, step)

    add_heading(doc, "14.3 Routing rules", 2)
    add_table(
        doc,
        [
            ["Question type", "Route"],
            ["Intent, dealer, consent, task", "Flow actions only"],
            ["Talking points, specs, playbook, objections", "Answer Question with Knowledge"],
            ["Discount / financing approval", "Blocked"],
        ],
    )

    add_heading(doc, "15. Streaming Ingestion API (live demo)", 1)
    add_para(
        doc,
        "Replace static CSV engagement with near-real-time events. Karan Singh (LEAD-2002) starts MEDIUM; "
        "post TestRideRequested via Ingestion API → intent becomes HIGH after ~3–5 minutes.",
    )
    add_heading(doc, "15.1 Prerequisites", 2)
    for item in [
        "Ingestion API connector SVT_Engagement_Events + schema YAML uploaded",
        "Data stream: Ingestion API → engagement_event → Website Engagement DMO",
        "External Client App with cdp_ingest_api, cdp_api, api scopes",
        "PowerShell script: docs/scripts/ingest-svt-engagement-event.ps1",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "15.2 Live demo steps", 2)
    add_table(
        doc,
        [
            ["Step", "Action", "Expected"],
            ["1", "Karan Singh — intent and dealer?", "MEDIUM, SVT Pune West"],
            ["2", "POST TestRideRequested for LEAD-2002", "HTTP 202"],
            ["3", "Wait 3–5 min", "New row in Website Engagement"],
            ["4", "Re-ask Karan's intent", "HIGH"],
        ],
    )

    add_heading(doc, "16. Phase 5 build roadmap", 1)
    add_table(
        doc,
        [
            ["Item", "Description", "Priority"],
            ["5a", "Complete Data Library + instructions", "Now"],
            ["5b", "Wire Data Library to Rider subagent", "High"],
            ["5c", "SVT Get Prospect Briefing unified Flow", "High"],
            ["5d", "Live Karan ingest rehearsal", "High"],
            ["5e", "Product catalog DMO + Get Model Specs", "Optional"],
            ["5f", "SVT Advisor Console Lightning page", "Optional"],
        ],
    )

    add_heading(doc, "17. Preview test questions", 1)
    add_heading(doc, "17.1 Flows only", 2)
    add_table(
        doc,
        [
            ["#", "Question", "Expected"],
            ["1", "Isha Sharma — intent and dealer?", "HIGH, Bengaluru Central"],
            ["2", "Karan Singh — intent and dealer?", "MEDIUM, Pune West"],
            ["3", "Aditi Menon — intent and dealer?", "MEDIUM, consent false, no task"],
            ["4", "Rohit Verma engagement level?", "LOW, Delhi South"],
        ],
    )

    add_heading(doc, "17.2 Data Library", 2)
    add_table(
        doc,
        [
            ["#", "Question", "Expected"],
            ["5", "Test ride steps before and after?", "Cites Playbook"],
            ["6", "Highlight Stride 200 for Naveen?", "Cites One-Pager, 3.2 kWh"],
            ["7", "Motor peak power Stride 200?", "4.2 kW from One-Pager"],
            ["8", "Handle too expensive objection?", "Cites Objection Guide"],
        ],
    )

    add_heading(doc, "17.3 Combined + guardrails", 2)
    add_table(
        doc,
        [
            ["#", "Question", "Expected"],
            ["9", "Naveen briefing + talking points", "Flow + Library"],
            ["10", "Yes, create test ride task for Naveen", "CREATED + emailSent"],
            ["11", "Create task for Aditi", "Blocked — consent"],
            ["12", "Approve 15% discount", "Refused"],
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Appendix A — Lead Engagement subagent instructions (full)", 1)
    for line in LEAD_INSTRUCTIONS.splitlines():
        if line.strip():
            add_para(doc, line)
        else:
            doc.add_paragraph()

    add_heading(doc, "Appendix B — Repo reference files", 1)
    add_table(
        doc,
        [
            ["File", "Purpose"],
            ["docs/svt-data-cloud-agentforce-demo.md", "Main architecture guide"],
            ["docs/svt-data-library-org-setup.md", "Data Library step-by-step"],
            ["docs/svt-agentforce-phase5-build.md", "Phase 5 build plan"],
            ["docs/svt-realtime-ingestion-guide.md", "Streaming Ingestion API"],
            ["docs/svt-rag-data-library/*.pdf", "Advisor PDF sources"],
        ],
    )


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title_page(doc)
    md_to_docx(doc, DOCS / "svt-data-cloud-agentforce-demo.md")
    add_supplement_sections(doc)
    set_update_fields(doc)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
